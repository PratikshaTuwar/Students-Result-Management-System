from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
import os

FILE_NAME = "student_results.xlsx"


# ----------------------------------------
# Create Excel file if it does not exist
# ----------------------------------------
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No",
            "Name",
            "Class",
            "Subject 1",
            "Subject 2",
            "Subject 3",
            "Subject 4",
            "Subject 5",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ]

        sheet.append(headers)

        # Make headings bold
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        workbook.save(FILE_NAME)


# ----------------------------------------
# Calculate result
# ----------------------------------------
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    # Pass only if every subject has 35 or more
    if any(mark < 35 for mark in marks):
        grade = "F"
        status = "FAIL"
    elif percentage >= 90:
        grade = "A+"
        status = "PASS"
    elif percentage >= 80:
        grade = "A"
        status = "PASS"
    elif percentage >= 70:
        grade = "B"
        status = "PASS"
    elif percentage >= 60:
        grade = "C"
        status = "PASS"
    elif percentage >= 50:
        grade = "D"
        status = "PASS"
    elif percentage >= 35:
        grade = "E"
        status = "PASS"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


# ----------------------------------------
# Add student
# ----------------------------------------
def add_student():
    print("\n========================================")
    print("          ADD STUDENT RESULT")
    print("========================================")

    # Roll number
    while True:
        try:
            roll_no = int(input("Enter Roll No: "))

            if roll_no <= 0:
                print("Roll No must be greater than 0.")
                continue

            break

        except ValueError:
            print("Please enter a valid number.")

    # Check duplicate roll number
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == roll_no:
            print("A student with this Roll No already exists.")
            workbook.close()
            return

    workbook.close()

    # Student name
    while True:
        name = input("Enter Student Name: ").strip()

        if name:
            break

        print("Name cannot be empty.")

    # Class / Course
    while True:
        course = input("Enter Course/Class: ").strip()

        if course:
            break

        print("Course/Class cannot be empty.")

    # Enter marks
    marks = []

    for i in range(1, 6):
        while True:
            try:
                mark = float(input(f"Enter Marks of Subject {i}: "))

                if mark < 0 or mark > 100:
                    print("Marks must be between 0 and 100.")
                    continue

                marks.append(mark)
                break

            except ValueError:
                print("Please enter a valid number.")

    # Calculate result
    total, percentage, grade, status = calculate_result(marks)

    # Save to Excel
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    sheet.append([
        roll_no,
        name,
        course,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("\n----------------------------------------")
    print("       RESULT SAVED SUCCESSFULLY")
    print("----------------------------------------")
    print(f"Roll No     : {roll_no}")
    print(f"Name        : {name}")
    print(f"Class       : {course}")
    print(f"Total       : {total}")
    print(f"Percentage  : {percentage:.2f}%")
    print(f"Grade       : {grade}")
    print(f"Status      : {status}")
    print("----------------------------------------")


# ----------------------------------------
# Get individual student result
# ----------------------------------------
def get_result():
    print("\n========================================")
    print("          GET STUDENT RESULT")
    print("========================================")

    while True:
        try:
            roll_no = int(input("Enter Roll No: "))
            break

        except ValueError:
            print("Please enter a valid Roll No.")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] == roll_no:
            found = True

            print("\n--------------------------------")
            print("         STUDENT RESULT")
            print("--------------------------------")
            print(f"Roll No     : {row[0]}")
            print(f"Name        : {row[1]}")
            print(f"Class       : {row[2]}")
            print(f"Total       : {row[8]}")
            print(f"Percentage  : {row[9]:.2f}%")
            print(f"Grade       : {row[10]}")
            print(f"Status      : {row[11]}")
            print("--------------------------------")

            break

    workbook.close()

    if not found:
        print("\nStudent with this Roll No was not found.")


# ----------------------------------------
# Show all student data
# ----------------------------------------
def show_all_data():
    print("\n========================================")
    print("          ALL STUDENT DATA")
    print("========================================")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    if sheet.max_row <= 1:
        print("No student records found.")
        workbook.close()
        return

    # Table headings
    print(
        f"{'Roll':<8}"
        f"{'Name':<18}"
        f"{'Class':<12}"
        f"{'Total':<10}"
        f"{'Percent':<12}"
        f"{'Grade':<8}"
        f"{'Status':<8}"
    )

    print("-" * 76)

    for row in sheet.iter_rows(min_row=2, values_only=True):

        percentage = float(row[9])

        print(
            f"{str(row[0]):<8}"
            f"{str(row[1])[:17]:<18}"
            f"{str(row[2])[:11]:<12}"
            f"{str(row[8]):<10}"
            f"{percentage:.2f}%     "
            f"{str(row[10]):<8}"
            f"{str(row[11]):<8}"
        )

    print("-" * 76)

    workbook.close()


# ----------------------------------------
# Main menu
# ----------------------------------------
def menu():
    create_excel_file()

    while True:

        print("\n")
        print("========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ").strip()

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nThank you for using Student Result Management System!")
            print("Program closed.")
            break

        else:
            print("\nInvalid choice!")
            print("Please enter a number from 1 to 4.")


# ----------------------------------------
# Start program
# ----------------------------------------
if __name__ == "__main__":
    menu()